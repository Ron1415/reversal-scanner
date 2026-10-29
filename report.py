"""Writes scan results to CSV, Excel and the web dashboard (site/index.html)."""
from __future__ import annotations

import datetime as dt
import json
import os
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TFS = ("daily", "weekly", "monthly")
COLS = ["ticker", "direction", "stage", "idea", "score", "count", "close", "change_pct", "entry", "stop",
        "risk_pct", "moved_pct", "rsi", "factors", "bar_date", "chart"]
STAGE_NAMES = {"forming": "Tail forming now", "day2": "Tail 1-2 days ago", "watch": "Watch: into a level"}


def _table(rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    if df.empty:
        return pd.DataFrame(columns=COLS)
    df["factors"] = df["factors"].apply(" | ".join)
    df["stage"] = [STAGE_NAMES.get(s, "No longer valid") if q else "No longer valid"
                   for s, q in zip(df["stage"], df["qualifies"])]
    for col in COLS:
        if col not in df:
            df[col] = None
    df["chart"] = "https://www.tradingview.com/chart/?symbol=" + df["ticker"]
    return df[COLS]


def _write_xlsx(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    red, green = PatternFill("solid", fgColor="FBE3E1"), PatternFill("solid", fgColor="E2F3E6")
    status_fill = {"Tail forming now": "16A34A", "Tail 1-2 days ago": "1D4ED8", "Watch: into a level": "F59E0B",
                   "No longer valid": "9CA3AF"}
    head = PatternFill("solid", fgColor="1F2937")
    widths = dict(ticker=9, direction=10, stage=20, idea=70, score=7, count=7, close=10, change_pct=10, entry=10,
                  stop=10, risk_pct=9, moved_pct=10, rsi=6, factors=90, bar_date=11, chart=12)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name.capitalize(), index=False)
            ws = xw.sheets[name.capitalize()]
            ws.freeze_panes = "B2"
            for j, col in enumerate(df.columns, 1):
                cell = ws.cell(row=1, column=j)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
                ws.column_dimensions[get_column_letter(j)].width = widths.get(col, 12)
            for i, r in enumerate(df.itertuples(index=False), 2):
                fill = red if r.direction == "bearish" else green
                for j in range(1, len(df.columns) + 1):
                    ws.cell(row=i, column=j).fill = fill
                st = ws.cell(row=i, column=COLS.index("stage") + 1)
                st.fill = PatternFill("solid", fgColor=status_fill.get(r.stage, "9CA3AF"))
                st.font = Font(bold=True, color="FFFFFF")
                ws.cell(row=i, column=COLS.index("idea") + 1).alignment = Alignment(wrap_text=True)
                ws.cell(row=i, column=COLS.index("factors") + 1).alignment = Alignment(wrap_text=True)
                link = ws.cell(row=i, column=COLS.index("chart") + 1)
                link.hyperlink, link.value, link.font = r.chart, "Open chart", Font(color="1D4ED8", underline="single")
            ws.auto_filter.ref = ws.dimensions


def write_outputs(tf: str, rows: list[dict], now: dt.datetime, root: Path, mode: str = "full") -> None:
    stamp = now.strftime("%Y-%m-%d")
    hist = root / "history" / tf
    hist.mkdir(parents=True, exist_ok=True)
    data_dir = root / "site" / "data"
    data_dir.mkdir(parents=True, exist_ok=True)

    # 1) History archive (one CSV per run)
    _table(rows).to_csv(hist / f"{stamp}_{now.strftime('%H%M')}.csv", index=False)

    # 2) Latest JSON for this timeframe
    payload = dict(timeframe=tf, generated=now.strftime("%Y-%m-%d %H:%M %Z"),
                   generated_iso=now.isoformat(), mode=mode, rows=rows)
    (data_dir / f"{tf}.json").write_text(json.dumps(payload))

    # 3) Rebuild the combined workbook and dashboard from all latest JSON files
    latest = {}
    for name in TFS:
        f = data_dir / f"{name}.json"
        latest[name] = json.loads(f.read_text()) if f.exists() else dict(timeframe=name, generated=None, rows=[])
    _write_xlsx(root / "site" / "reversal_signals.xlsx", {k: _table(v["rows"]) for k, v in latest.items()})

    write_pages(root, latest)


def _cfg() -> dict:
    # e.g. GITHUB_WORKFLOW_REF = "user/repo/.github/workflows/scan.yml@refs/heads/main"
    wf = os.environ.get("GITHUB_WORKFLOW_REF", "").split("@")[0].rsplit("/", 1)[-1] or "scan.yml"
    return dict(repo=os.environ.get("GITHUB_REPOSITORY", ""), branch=os.environ.get("GITHUB_REF_NAME", "main"),
                workflow=wf)


def write_pages(root: Path, latest: dict | None = None) -> None:
    """(Re)build site/index.html, site/tracker.html and site/store.js."""
    site = root / "site"
    (site / "data").mkdir(parents=True, exist_ok=True)
    if latest is None:
        latest = {}
        for name in TFS:
            f = site / "data" / f"{name}.json"
            latest[name] = json.loads(f.read_text()) if f.exists() else dict(timeframe=name, generated=None, rows=[])
    cfg = json.dumps(_cfg())
    tpl = (root / "dashboard_template.html").read_text()
    (site / "index.html").write_text(tpl.replace("/*__DATA__*/null", json.dumps(latest)).replace("/*__CFG__*/null", cfg))
    tpl = (root / "tracker_template.html").read_text()
    (site / "tracker.html").write_text(tpl.replace("/*__CFG__*/null", cfg))
    (site / "store.js").write_text((root / "store.js").read_text())
